import csv
from datetime import date, datetime
from pathlib import Path
from docx import Document
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer,
)
from reportlab.lib.styles import StyleSheet1, getSampleStyleSheet
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from classes import calc_age
from operations import search, get_relatives_by_patient_id, ambiguity_check, resolve_selection


def build_patient_record(search_field:str, search_value:str, row_identifier:str | None = None) -> tuple[str, dict | list | None]:
    result:dict[str, str] | list[dict] | None = search("users", search_field, search_value)
    if result is None:
        print(f"The value '{search_value}' could not be found in the field '{search_field}' of the 'users' table")        
        return ("not_found", None)

    ambiguity:str | None = ambiguity_check("users", search_field, result)
    chosen:dict | None = {}

    if ambiguity == "Ambiguous":
        if not isinstance(result, list):
            print("Invalid data format.")
            return ("failed", None)
        if row_identifier is None:
            print("Multiple matches found. Please Select any one by their 'patient_id'")
            return ("ambiguous", result)
        else:
            chosen = resolve_selection("users", result, row_identifier)
            if chosen is None:
                return ("not_found", None)
    elif ambiguity == "Unambiguous":
        chosen = result[0] if isinstance(result, list) else result
    else:
        print(f"The value '{search_value}' could not be found in the field '{search_field}' of the 'users' table")
        return ("not_found", None)

    dob:date = datetime.strptime(chosen["date_of_birth"], "%d-%m-%Y").date()
    age:int = calc_age(date.today(), dob)
    relatives:list[dict[str, str]] | None = get_relatives_by_patient_id(chosen["patient_id"])
    if relatives is None:
        relatives = []
    record:dict[str, dict | list] = {"user":{**chosen, "age": age}, "relatives":relatives}
    return ("success", record)


def prompt_export_path() -> Path | None:
    while True:
        path:str = input("Enter the absolute path to where you would like to export the file:").strip()
        if path:
            if Path(path).is_absolute():
                return Path(path)
            else:
                print("Incorrect Absolute Path Entered.")
                re_prompt:str = input("[r]e-enter / [d]efault").strip().lower()
                if re_prompt == "r":
                    continue
                elif re_prompt == "d":
                    return None
                else:
                    print("Invalid Choice.")
                    continue
        else:
            return None


def resolve_export_path(filename:str, path:str | Path | None = None) -> Path:
    export_to_path:Path = Path(__file__).resolve().parent.parent / "exports" / filename   
    if path:
        if Path(path).is_absolute():
            export_to_path = Path(path) / filename
    export_to_path.parent.mkdir(parents=True, exist_ok=True)
    return export_to_path


def format_time(select:str) -> str | None:
    timestamp:datetime = datetime.now()
    if select == "file_name":
        return timestamp.strftime("%d-%m-%Y_%H-%M-%S")
    elif select == "in_file":
        return timestamp.strftime("%d-%m-%Y %H:%M:%S")
    return None

# ----------------- SHARED EXPORT HELPERS -----------------

def write_txt_record(file, record: dict) -> None:
    file.write("----------User Data----------\n")
    for key, val in record['user'].items():
        file.write(f"{key}: {val}\n")
    
    file.write("----------Relatives Data----------\n")
    if record['relatives']:
        for i, relative in enumerate(record['relatives']):
            file.write(f"Relative{str(i)}\n")
            for key, val in relative.items():
                val = val or ""  # Guard against explicit None for middle_name
                file.write(f"{key}: {val}\n")
    else:
        file.write("This User doesn't yet have any Relatives attached.\n")


def format_csv_row(record: dict) -> dict:
    relatives_str = "; ".join(
        f"{rel['first_name']} {rel.get('middle_name') or ''} {rel['last_name']}({rel['email']}, {rel['phone_number']})"
        for rel in record['relatives']
    )
    # Replaces multiple contiguous spaces with a single one (e.g. if middle_name was empty)
    relatives_str = " ".join(relatives_str.split())
    return {**record['user'], "relatives": relatives_str}


def write_docx_record(doc, record: dict, is_admin: bool = False) -> None:
    doc.add_heading("User Details" if not is_admin else "User Data", level=2 if not is_admin else 3)
    for key, val in record['user'].items():
        doc.add_paragraph(f"{key}: {val}")
    
    doc.add_heading("Relatives Details" if not is_admin else "Relatives Data", level=2 if not is_admin else 3)
    if record['relatives']:
        for i, relative in enumerate(record['relatives']):
            if is_admin:
                doc.add_heading(f"Relative{str(i)}", level=4)
            else:
                doc.add_paragraph(f"Relative{str(i)}")
            for key, val in relative.items():
                val = val or ""
                doc.add_paragraph(f"{key}: {val}")
    else:
        doc.add_paragraph("This User doesn't yet have any Relatives attached.")


def build_pdf_story(story: list, record: dict, styles: StyleSheet1) -> None:
    story.append(Paragraph("User Data" if "Patient" in story[-1].text else "Patient Details", styles['Heading2' if "Patient Details" in story[-1].text else 'Heading3']))
    for key, val in record['user'].items():
        story.append(Paragraph(f"{key}: {val}", styles['Normal']))

    story.append(Paragraph("Relatives", styles['Heading2'] if "Patient Details" in story[-1].text else styles['Heading3']))
    if record['relatives']:
        header = [Paragraph(h, styles["Normal"]) for h in ["First Name", "Middle Name", "Last Name", "Email", "Phone Number"]]
        rows = [
            [
                Paragraph(rel["first_name"], styles["Normal"]),
                Paragraph(rel.get("middle_name") or "", styles["Normal"]),
                Paragraph(rel["last_name"], styles["Normal"]),
                Paragraph(rel["email"], styles["Normal"]),
                Paragraph(rel["phone_number"], styles["Normal"]),
            ]
            for rel in record['relatives']
        ]
        table = Table([header, *rows], colWidths=[70, 70, 80, 170, 80])
        table.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
            ("BACKGROUND", (0, 0), (-1, 0), colors.gray),
        ]))        
        story.append(table)
    else:
        story.append(Paragraph("This User doesn't yet have any Relatives attached", styles['Normal']))

# ----------------- PATIENT EXPORTS -----------------

def export_patient_txt(record: dict, path: str | Path | None = None) -> str:
    patient_id: str = record['user']["patient_id"]
    write_to_path = resolve_export_path(f"{patient_id}_{format_time('file_name')}.txt", path)
    
    with open(write_to_path, "w", encoding="utf-8") as file:
        file.write(f"Export Time: {format_time('in_file')}\n")
        write_txt_record(file, record)
        
    return str(write_to_path)


def export_patient_csv(record: dict, path: str | None = None) -> str:
    patient_id: str = record['user']["patient_id"]
    write_to_path = resolve_export_path(f"{patient_id}_{format_time('file_name')}.csv", path)
    
    row = format_csv_row(record)
    with open(write_to_path, "w", newline="", encoding="utf-8") as file:
        print("Note: each patient's relatives are stored in a single ';'-separated field.")
        writer = csv.DictWriter(file, fieldnames=list(row.keys()))
        writer.writeheader()
        writer.writerow(row)
        
    return str(write_to_path)


def export_patient_docx(record: dict, path: str | None = None) -> str:
    patient_id: str = record['user']['patient_id']
    write_to_path = resolve_export_path(f"{patient_id}_{format_time('file_name')}.docx", path)
    
    doc = Document()
    doc.add_heading("Patient Record", level=1)
    doc.add_paragraph(f"Export Time: {format_time('in_file')}")
    write_docx_record(doc, record)
    doc.save(str(write_to_path))
    
    return str(write_to_path)


def export_patient_pdf(record: dict, path: str | Path | None = None) -> str:
    patient_id: str = record['user']['patient_id']
    write_to_path = resolve_export_path(f"{patient_id}_{format_time('file_name')}.pdf", path)
    
    styles: StyleSheet1 = getSampleStyleSheet()
    story: list = [
        Paragraph("Patient Record", styles["Title"]),
        Paragraph(f"Export Time: {format_time('in_file')}", styles["Normal"]),
        Paragraph("Patient Details", styles["Title"]) # Hack to make the builder work
    ]
    build_pdf_story(story, record, styles)
    
    doc = SimpleDocTemplate(str(write_to_path), pagesize=A4)
    doc.build(story)
    
    return str(write_to_path)

# ----------------- ADMIN EXPORTS -----------------

def build_admin_records(patient_id_list: list[str]) -> list[dict]:
    data_list: list[dict] = []
    for pid in patient_id_list:
        record: tuple = build_patient_record("patient_id", pid)
        if record[0] == "success":
            data_list.append(record[1])
        else:
            print(f"The patient_id {pid} does not exist")
    return data_list


def export_admin_txt(records: list[dict], path: str | Path | None = None) -> str: 
    write_to_path = resolve_export_path(f"admin_export_{format_time('file_name')}.txt", path)
    
    with open(write_to_path, "w", encoding="utf-8") as file:
        file.write(f"Export Time: {format_time('in_file')}")
        for record in records:
            file.write(f"\n=========={record['user']['patient_id']}==========\n")
            write_txt_record(file, record)
            
    return str(write_to_path)


def export_admin_csv(records: list[dict], path: str | Path | None = None) -> str:
    write_to_path = resolve_export_path(f"admin_export_{format_time('file_name')}.csv", path)
    
    fieldnames = ["patient_id", "first_name", "middle_name", "last_name", "date_of_birth",
                 "age", "symptoms", "email", "phone_number", "relatives"]
                 
    with open(write_to_path, "w", newline="", encoding="utf-8") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames)
        writer.writeheader()
        print("Note: each patient's relatives are stored in a single ';'-separated field.")
        for record in records:
            writer.writerow(format_csv_row(record))
            
    return str(write_to_path)


def export_admin_docx(records: list[dict], path: str | Path | None = None) -> str:
    write_to_path = resolve_export_path(f"admin_export_{format_time('file_name')}.docx", path)
    
    doc = Document()
    doc.add_heading("Admin Export", level=1)
    doc.add_paragraph(f"Export Time: {format_time('in_file')}")
    
    for record in records:
        doc.add_heading(f"Patient {record['user']['patient_id']}", level=2)
        write_docx_record(doc, record, is_admin=True)
        
    doc.save(str(write_to_path))
    return str(write_to_path)


def export_admin_pdf(records: list[dict], path: str | Path | None = None) -> str:
    write_to_path = resolve_export_path(f"admin_export_{format_time('file_name')}.pdf", path)
    
    styles: StyleSheet1 = getSampleStyleSheet()
    story: list = [
        Paragraph("Admin Export", styles['Title']),
        Paragraph(f"Export Time: {format_time('in_file')}", styles['Normal'])
    ]
    
    for record in records:
        story.append(Paragraph(f"Patient {record['user']['patient_id']}", styles['Heading2']))
        build_pdf_story(story, record, styles)
        
    doc = SimpleDocTemplate(str(write_to_path), pagesize=A4)
    doc.build(story)
    
    return str(write_to_path)
